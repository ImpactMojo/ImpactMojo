"""RAID Log (Excel workbook and Word guide)."""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

from xl_pack import grid_sheet, package, readme

BASE = "ImpactMojo-RAID-Log"
FILE = BASE + ".zip"
LEVEL = ["1 Low", "2 Medium", "3 High"]
STATUS = ["Open", "In progress", "Closed"]

SCORE = lambda r: f'=IF(OR(E{r}="",F{r}=""),"",VALUE(LEFT(E{r},1))*VALUE(LEFT(F{r},1)))'
BAND = lambda r: f'=IF(G{r}="","",IF(G{r}>=6,"Act now",IF(G{r}>=3,"Watch","Accept")))'


def build(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme(ws, "RAID log", "Risks, assumptions, issues and dependencies for an evaluation, in one place, with a score that tells you what to act on first.", [
        ("What is in the workbook", True),
        ("Risks: things that might go wrong. Chance and effect are each scored 1 to 3. The score is chance times effect, and the band reads Act now (6 or more), Watch (3 or 4) or Accept (1 or 2).", False),
        ("Assumptions: things you are relying on that are not yet confirmed, each with the date you will check it.", False),
        ("Issues: things that have already gone wrong, with an owner and a date.", False),
        ("Dependencies: things you need from someone else before you can go on, with the date you need them.", False),
        ("Worked example: a log for a maternal and newborn health evaluation in three blocks.", False),
        ("How to use it", True),
        ("Yellow cells are yours to fill. Grey cells calculate. Review the log at the start of every weekly call, and close an item only when you can say what happened.", False),
    ])
    risk_cols = [("ID", 7), ("Risk (if X happens, then Y)", 44), ("Cause", 28), ("Phase", 14), ("Chance", 11), ("Effect", 11), ("Score", 8), ("Band", 11), ("Action to reduce it", 40), ("Owner", 14), ("Review by", 12), ("Status", 12)]
    w = wb.create_sheet("Risks")
    grid_sheet(w, "Risks", "Score the chance and the effect from 1 to 3. The score and the band calculate.", risk_cols, 25,
               calc={6: SCORE, 7: BAND}, drops={4: LEVEL, 5: LEVEL, 11: STATUS}, fmts={10: "dd mmm yyyy"})
    w.conditional_formatting.add("H5:H29", FormulaRule(formula=['H5="Act now"'], fill=PatternFill("solid", bgColor="F8D7D3")))
    w.conditional_formatting.add("H5:H29", FormulaRule(formula=['H5="Watch"'], fill=PatternFill("solid", bgColor="FFF1C9")))
    a = wb.create_sheet("Assumptions")
    grid_sheet(a, "Assumptions", "Write what you are relying on, and when you will check it.", [("ID", 7), ("Assumption", 52), ("Why we believe it", 36), ("How to check", 34), ("Check by", 12), ("Owner", 14), ("Result", 16)], 20,
               drops={6: ["Not yet checked", "Holds", "Does not hold"]}, fmts={4: "dd mmm yyyy"})
    i = wb.create_sheet("Issues")
    grid_sheet(i, "Issues", "Problems that have already happened.", [("ID", 7), ("Issue", 48), ("Date raised", 12), ("Effect on the work", 36), ("Action", 38), ("Owner", 14), ("Due", 12), ("Status", 12)], 20,
               drops={7: STATUS}, fmts={2: "dd mmm yyyy", 6: "dd mmm yyyy"})
    d = wb.create_sheet("Dependencies")
    grid_sheet(d, "Dependencies", "What you need from others, and by when.", [("ID", 7), ("What we need", 44), ("From whom", 22), ("Needed by", 12), ("Why it matters", 36), ("Promised?", 12), ("Status", 12), ("Next step", 34)], 20,
               drops={5: ["Yes", "No", "Not yet asked"], 6: STATUS}, fmts={3: "dd mmm yyyy"})
    x = wb.create_sheet("Worked example")
    ex = [
        ["R1", "If the state does not release the facility list by week 2, then sampling slips and fieldwork starts late", "Approval sits with the mission director", "Inception", "2 Medium", "3 High", None, None, "Ask for the list in the kick-off; name a second contact", "Project manager", "2026-11-20", "Open"],
        ["R2", "If enumerators cannot reach two hill-block facilities in the monsoon, then the sample is short", "Roads close from June", "Fieldwork", "2 Medium", "2 Medium", None, None, "Schedule those blocks first; keep a replacement list", "Field lead", "2026-12-10", "Open"],
        ["R3", "If the client sends comments in several rounds, then the review runs past its fixed period", "Four reviewers, no single point of contact", "Reporting", "3 High", "2 Medium", None, None, "Ask for one consolidated set of comments in the statement of work", "Team leader", "2026-12-01", "In progress"],
        ["R4", "If a respondent is identified from a case study, then we breach consent", "Small facility, few staff", "Reporting", "1 Low", "3 High", None, None, "Remove identifying detail; read each case back to the respondent", "Team leader", "2027-02-10", "Open"],
    ]
    grid_sheet(x, "Worked example: risks", "Four risks for the Surakshit Maa evaluation. The score and band calculate.", risk_cols, 4, rows=ex,
               calc={6: SCORE, 7: BAND}, drops={4: LEVEL, 5: LEVEL, 11: STATUS}, fmts={10: "dd mmm yyyy"})
    x.title = "Worked example"
    return package(out, BASE, wb, GUIDE)


GUIDE = dict(
    file=BASE + "-GUIDE.docx",
    cover=("RAID Log", "A single log of risks, assumptions, issues and dependencies, scored so that the team acts on the right things first", "Guide to the Excel workbook · ₹199"),
    inside=[
        "**Part A.** Why a log beats a worry, and what RAID means.",
        "**Part B.** The method in five steps: list, score, assign, review and close.",
        "**Part C.** Reference assets: wording for a risk, a review routine, typical evaluation risks and a glossary.",
        "**Part D.** A worked example: four risks for a maternal and newborn health evaluation.",
    ],
    intro="This guide goes with the Excel workbook. Keep one log for each assignment and review it at every weekly call. A log that nobody opens is a list of old worries.",
    parts=[
        ("PART A", "Why keep a log", [
            ("body", "Evaluations run into the same trouble again and again: a list that arrives late, a monsoon, a reviewer who changes the questions in week ten. Most of it can be seen in advance. A log makes the team say it aloud, give it an owner and a date, and look at it every week."),
            ("body", "RAID stands for risks, assumptions, issues and dependencies. A risk may happen. An assumption is something you rely on without having confirmed it. An issue has already happened. A dependency is something you need from someone else. Keeping them on four sheets of one workbook stops the same worry being written down four times."),
            ("body", "Three rules keep the log useful:"),
            ("bullets", [
                "Write each risk as 'if X happens, then Y', so that cause and effect are both visible.",
                "Give every item one named owner and one date.",
                "Review the log at the start of the weekly call, and keep the call short by looking only at items that are Act now or due.",
            ]),
        ]),
        ("PART B", "The method", [
            ("h", "1.  List"),
            ("purpose", "Get the worries out of people's heads and into one place."),
            ("guide", "Run a 20-minute session at kick-off. Each person writes the things that could go wrong, the things they are assuming and the things they need from others. Sort them into the four sheets."),
            ("good", "Every risk reads 'if X, then Y', and every assumption has a way to check it."),
            ("mistakes", "Listing a worry in general terms, such as 'poor data quality', that nobody can act on."),
            ("h", "2.  Score"),
            ("purpose", "Decide which risks need action now."),
            ("guide", "Score the chance and the effect from 1 to 3 in the dropdowns. The sheet multiplies them and sets the band: Act now for 6 or more, Watch for 3 or 4, Accept for 1 or 2."),
            ("table", ["Score", "Band", "What to do"], [
                ["6 or 9", "Act now", "Agree an action this week and give it an owner."],
                ["3 or 4", "Watch", "Name a trigger that would move it up, and check it at each review."],
                ["1 or 2", "Accept", "Record it and leave it."],
            ], [2.6, 3.0, 10.3]),
            ("good", "Scores differ across items. A log in which everything is a 9 tells you nothing."),
            ("mistakes", "Scoring by the loudest voice in the room. Ask for the chance first and the effect second."),
            ("h", "3.  Assign"),
            ("purpose", "Make someone answerable."),
            ("guide", "Write one owner and one date for each item. The owner is the person who will raise it if it moves, whoever caused it."),
            ("good", "No item has two owners or none."),
            ("mistakes", "Writing 'team' in the owner column."),
            ("h", "4.  Review"),
            ("purpose", "Keep the log current."),
            ("guide", "Open the Risks sheet at the weekly call. Look at the Act now rows and any row past its review date. Change the chance or effect if the facts have moved, and add new items from the week."),
            ("good", "Each review changes at least one row."),
            ("mistakes", "Reading the log aloud without changing anything."),
            ("h", "5.  Close"),
            ("purpose", "Show what happened, so the next assignment starts smarter."),
            ("guide", "Close an item only when you can say in a line what happened. Carry the lessons into the after-action review."),
            ("good", "The closed items form the first draft of your lessons."),
            ("mistakes", "Deleting closed items."),
        ]),
        ("PART C", "Reference assets", [
            ("h", "C1.  Typical evaluation risks to test your own list against"),
            ("bullets", [
                "Access to data, facilities or respondents is delayed or refused.",
                "Weather, festivals or elections close the field window.",
                "A key team member leaves or is ill.",
                "The questions change after inception.",
                "Review comments arrive in several rounds, or late.",
                "A respondent can be identified from the report.",
                "The client's own data turn out to be incomplete.",
                "A subcontractor delivers late.",
            ]),
            ("h", "C2.  Glossary"),
            ("table", ["Term", "Meaning"], [
                ["Risk", "Something that might happen and would hurt the work."],
                ["Assumption", "Something you rely on that you have not confirmed."],
                ["Issue", "A problem that has already happened."],
                ["Dependency", "Something you need from someone else before you can go on."],
                ["Trigger", "An event that tells you a risk is becoming real."],
            ], [3.6, 12.3]),
        ]),
        ("PART D", "Worked example: maternal and newborn health evaluation", [
            ("table", ["ID", "Risk", "Chance", "Effect", "Score", "Band"], [
                ["R1", "State does not release the facility list by week 2", "2", "3", "6", "Act now"],
                ["R2", "Two hill-block facilities unreachable in the monsoon", "2", "2", "4", "Watch"],
                ["R3", "Several rounds of client comments", "3", "2", "6", "Act now"],
                ["R4", "A respondent identified from a case study", "1", "3", "3", "Watch"],
            ], [1.2, 8.4, 1.8, 1.8, 1.4, 1.3]),
            ("body", "R1 and R3 are Act now. For R1 the project manager asks for the list at the kick-off and names a second contact. For R3 the team leader adds a single consolidated set of comments to the statement of work. R4 is on watch, and the trigger is the first draft of the case studies."),
            ("figures", "The risks and scores are illustrative. Replace them with your own."),
        ]),
    ],
)

PAGE = dict(
    slug="raid-log-template",
    title="RAID Log",
    file=FILE,
    ext="XLSX",
    eyebrow="Template · XLSX",
    crumb="Template",
    meta="An Excel RAID log for evaluations: risks scored by chance and effect, assumptions, issues and dependencies, with a Word guide and a worked example. ₹199.",
    og="A scored log of risks, assumptions, issues and dependencies for an evaluation. ₹199.",
    lede="One log of risks, assumptions, issues and dependencies, scored so that the team acts on the right things first.",
    note="Editable XLSX workbook + a Word guide · delivered to your inbox within 24 hours of payment.",
    letter_h="See trouble before it arrives",
    letter=[
        "Evaluations run into the same trouble again and again: a facility list that arrives late, a monsoon, a reviewer who changes the questions in week ten. Most of it can be seen in advance, and a log makes the team say it aloud, give it an owner and a date, and look at it every week.",
        "This workbook has four sheets, one each for risks, assumptions, issues and dependencies. Risks are scored by chance and effect from 1 to 3, and the score sets a band: Act now, Watch or Accept. A Word guide gives a five-step method, a list of typical evaluation risks and a worked example.",
        'It works beside the <a href="/products/weekly-status-note-template/">Weekly Status Note</a>, which reports the log to the client, and the <a href="/products/after-action-review-template/">After-Action Review</a>, which turns closed items into lessons.',
    ],
    inside=[
        "Four sheets: risks, assumptions, issues and dependencies",
        "Risk scoring from 1 to 3 on chance and effect, with a band that calculates",
        "Owner, review date and status on every item",
        "A Word guide with a five-step method and typical evaluation risks",
        "A worked example: four risks for a maternal and newborn health evaluation",
    ],
    sample=[
        ("Risks", "chance times effect, banded Act now, Watch or Accept"),
        ("Assumptions", "each with the date you will check it"),
        ("Issues", "owner, due date and status"),
        ("Dependencies", "what you need, from whom, by when"),
    ],
    sample_file=None,
    faq=[
        ("What format is it?", "An Excel workbook (.xlsx) with live formulas and dropdowns, plus a Word guide. It also opens in Google Sheets and LibreOffice."),
        ("How is the score worked out?", "The chance (1 to 3) is multiplied by the effect (1 to 3). A score of 6 or more is Act now, 3 or 4 is Watch and 1 or 2 is Accept."),
        ("Can I use it for more than one assignment?", "Yes. Save a copy for each assignment so that each log carries its own history."),
        ("How fast is delivery?", "Within 24 hours of confirming payment, usually much sooner."),
    ],
    search_desc="An Excel RAID log for evaluations with scored risks, assumptions, issues and dependencies, a Word guide and a worked example. Buy for ₹199 via UPI.",
    search_tags=["template", "RAID log", "risk", "project management", "consulting", "evaluation", "Excel", "premium", "paid", "₹199"],
    search_id="PROD-RAID-LOG",
    card="RAID Log",
    card_p="Risks, assumptions, issues and dependencies, scored.",
    before="Stakeholder Map",
)
