#!/usr/bin/env python3
"""Build the Evaluator's Desk deliverables and lint them against the owner's writing rules.

    python3 scripts/evaluator-pack/build.py OUT_DIR

Writes, into OUT_DIR:
  ImpactMojo-Bid-No-Bid-Scorecard.zip          (xlsx + Word guide)
  ImpactMojo-SOW-Change-Request-Template.docx
  ImpactMojo-Inception-Report-Template.docx

Exits 1 if the lint finds a banned word, an em dash, an emoji or a "not X" contrast.
Delivered files are not committed: the products live in the private `products`
bucket and in the Drive folder "ImpactMojo - Product Masters".
"""
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import inception  # noqa: E402
import scorecard  # noqa: E402
import sow  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

BANNED_WORDS = """delve utilise utilize leverage robust streamline harness foster underscore enhance testament pivotal
intricate intricacies crucial transformative groundbreaking tapestry landscape paradigm synergy architecture liminal
palpable ineffable visceral register ledger showcase meticulous meticulously notable notably noteworthy commendable
versatile versatility seamless seamlessly multifaceted nuanced holistic invaluable unparalleled garner bolster
elucidate interplay underpin underpinnings solidify vibrant bustling nestled picturesque breathtaking stunning boasts
renowned indelible poignant captivating evocative embark navigate unlock unleash empower elevate amplify resonate
evoke captivate facilitate spearhead realm domain arena myriad plethora sphere genuine genuinely honestly truly
arguably certainly comprehensive""".split()
# "comprehensive", "significant", "substantial" are allowed only with a concrete referent, so the first is flagged for a look.

PHRASES = [
    r"\bnot (?:only|just)\b", r",\s*not\b", r"\bit['’]s not\b", r"\bisn['’]t\b", r"\bdoesn['’]t\b",
    r"\brather than\b", r"\bthat said\b", r"\bmoreover\b", r"\bfurthermore\b", r"\badditionally\b",
    r"\bimportantly\b", r"\bin conclusion\b", r"\bultimately\b", r"\bkey takeaways?\b", r"\bneedless to say\b",
    r"\bit['’]s (?:important|worth) (?:to )?note\b", r"\bwhen it comes to\b", r"\bshed light\b", r"\bpave the way\b",
    r"\bin today['’]s\b", r"\bat the end of the day\b", r"\bhere['’]s the\b",
    r"—",  # em dash
    "[\U0001F300-\U0001FAFF☀-⛿✀-➿]",  # emoji and dingbats
]


def doc_text(path):
    d = Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                out.append(c.text)
    for s in d.sections:
        out += [p.text for p in s.footer.paragraphs]
    return out


def xlsx_text(path):
    wb = load_workbook(path)
    out = []
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith("="):
                    out.append(c.value)
                elif isinstance(c.value, str):
                    # formulas carry display text inside quotes
                    out += re.findall(r'"([^"]{12,})"', c.value)
    return out


def lint(name, lines):
    bad = []
    for line in lines:
        low = line.lower()
        for w in BANNED_WORDS:
            if re.search(r"\b" + w + r"\b", low):
                bad.append((name, "word: " + w, line[:110]))
        for pat in PHRASES:
            if re.search(pat, low if "u2014" not in pat else line):
                bad.append((name, "pattern: " + pat, line[:110]))
    return bad


def main(out):
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)
    sow.build(out / sow.FILE)
    inception.build(out / inception.FILE)
    scorecard.build_xlsx(out / scorecard.FILE_XLSX)
    scorecard.build_guide(out / scorecard.FILE_GUIDE)
    with zipfile.ZipFile(out / scorecard.FILE_ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(out / scorecard.FILE_XLSX, scorecard.FILE_XLSX)
        z.write(out / scorecard.FILE_GUIDE, scorecard.FILE_GUIDE)

    problems = []
    for f in (sow.FILE, inception.FILE, scorecard.FILE_GUIDE):
        problems += lint(f, doc_text(out / f))
    problems += lint(scorecard.FILE_XLSX, xlsx_text(out / scorecard.FILE_XLSX))
    legal = 0
    for f in (sow.FILE, inception.FILE, scorecard.FILE_GUIDE):
        legal += sum(1 for t in doc_text(out / f) if "[LEGAL CHECK" in t)
    for p in problems:
        print("LINT", *p, sep=" | ")
    print(f"built 4 files into {out}; {legal} [LEGAL CHECK] markers to clear before release; {len(problems)} lint hits")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "build"))
