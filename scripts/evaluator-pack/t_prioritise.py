"""Impact and Ease Prioritisation (Excel workbook and Word guide)."""
from openpyxl import Workbook

from xl_pack import grid_sheet, package, readme

BASE = "ImpactMojo-Impact-Ease-Prioritisation"
FILE = BASE + ".zip"
S5 = ["1", "2", "3", "4", "5"]

SCORE = lambda r: f'=IF(OR(C{r}="",D{r}="",E{r}=""),"",ROUND(C{r}*0.4+D{r}*0.35+E{r}*0.25,2))'
RANK = lambda r: f'=IF(F{r}="","",RANK(F{r},$F$5:$F$34))'
GRP = lambda r: f'=IF(OR(C{r}="",D{r}=""),"",IF(AND(C{r}>=4,D{r}>=4),"Do first",IF(C{r}>=4,"Plan it",IF(D{r}>=4,"Quick win","Park"))))'


def build(out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme(ws, "Impact and ease prioritisation", "A scoring sheet for choosing which recommendations to push first, and a grid that shows where each one sits.", [
        ("What is in the workbook", True),
        ("Score sheet: one row per recommendation, scored 1 to 5 on impact, ease and evidence. The weighted score is impact 40%, ease 35% and evidence 25%. The rank and the group calculate.", False),
        ("Groups: Do first (impact and ease both 4 or 5), Plan it (high impact, harder), Quick win (high ease, lower impact) and Park (the rest).", False),
        ("Grid: each recommendation placed by impact and ease, so a client can see the choice on one page.", False),
        ("Worked example: eight recommendations from a maternal and newborn health evaluation.", False),
        ("How to use it", True),
        ("Yellow cells are yours to fill. Grey cells calculate. Score with two or three people from the client and the team, argue about every 5 and every 1, and keep a note of why.", False),
        ("The weights are a starting point. Change them in the Weights sheet if the client has a different view, and say so in the report.", False),
    ])
    wt = wb.create_sheet("Weights")
    grid_sheet(wt, "Weights", "Change the weights if the client's priorities differ. They must add to 100%.", [("Criterion", 30), ("Weight", 12), ("What a 5 means", 60)], 3,
               rows=[["Impact", 0.4, "Fixes the problem the evaluation found, for most of the people affected."], ["Ease", 0.35, "Can be done within a year with the money, staff and authority the client has."], ["Evidence", 0.25, "Backed by more than one source in the evaluation."]], fmts={1: "0%"})
    wt["A9"] = "Total"
    wt["B9"] = "=SUM(B5:B7)"
    wt["B9"].number_format = "0%"
    cols = [("No.", 6), ("Recommendation", 48), ("Impact (1-5)", 11), ("Ease (1-5)", 11), ("Evidence (1-5)", 11), ("Weighted score", 11), ("Rank", 8), ("Group", 12), ("Why we scored it so", 44), ("Owner", 16)]
    calc = {5: lambda r: f'=IF(OR(C{r}="",D{r}="",E{r}=""),"",ROUND(C{r}*Weights!$B$5+D{r}*Weights!$B$6+E{r}*Weights!$B$7,2))', 6: RANK, 7: GRP}
    s = wb.create_sheet("Score sheet")
    grid_sheet(s, "Score sheet", "Score each recommendation from 1 to 5. The weighted score, rank and group calculate.", cols, 30, calc=calc, drops={2: S5, 3: S5, 4: S5})
    g = wb.create_sheet("Grid")
    from xl_pack import title, box, SOFT, NAVY
    from openpyxl.styles import Alignment, Font
    title(g, "Grid", "Recommendations placed by impact (rows) and ease (columns). Numbers refer to the Score sheet.")
    g.column_dimensions["A"].width = 14
    for c in "BCDEF":
        g.column_dimensions[c].width = 22
    g["A4"] = "Impact \\ Ease"
    for k in range(1, 6):
        g.cell(row=4, column=1 + k, value=f"Ease {k}").font = Font(name="Calibri", bold=True, color=NAVY)
        g.cell(row=4 + k, column=1, value=f"Impact {6 - k}").font = Font(name="Calibri", bold=True, color=NAVY)
    for ri in range(1, 6):
        imp = 6 - ri
        for ci in range(1, 6):
            parts = "&".join(f'IF(AND(\'Score sheet\'!$C${5 + n}={imp},\'Score sheet\'!$D${5 + n}={ci}),\'Score sheet\'!$A${5 + n}&" ","")' for n in range(30))
            cc = g.cell(row=4 + ri, column=1 + ci, value=f"={parts}")
            box(cc, SOFT, align="center")
            g.row_dimensions[4 + ri].height = 38
    g.sheet_view.showGridLines = False
    ex = [
        [1, "Give every health worker the new birth-asphyxia protocol", 5, 4, 5, None, None, None, "Largest gap in the findings; training budget exists", "State training lead"],
        [2, "Stock oxytocin at every sub-centre", 4, 3, 4, None, None, None, "Cold chain gaps in 9 of 30 sub-centres", "Supply chain"],
        [3, "Post staff rosters at each facility", 3, 5, 3, None, None, None, "Free and quick, but few effects alone", "Block medical officer"],
        [4, "Pay incentive on the day of delivery", 4, 2, 3, None, None, None, "Needs a finance order", "Mission director"],
        [5, "Digitise the mother-and-child records", 4, 2, 4, None, None, None, "Two years of work; strong findings", "IT cell"],
        [6, "Hold monthly review meetings by block", 3, 4, 3, None, None, None, "Meetings exist; the agenda does not", "District office"],
        [7, "Replace delivery tables", 2, 2, 2, None, None, None, "Cost high; little link to outcomes", "Procurement"],
        [8, "Run a mass campaign on danger signs", 3, 3, 2, None, None, None, "Weak evidence of effect", "Communications"],
    ]
    x = wb.create_sheet("Worked example")
    grid_sheet(x, "Worked example: eight recommendations", "From the Surakshit Maa evaluation. The score, rank and group calculate.", cols, 8, rows=ex, calc=calc, drops={2: S5, 3: S5, 4: S5})
    x.title = "Worked example"
    return package(out, BASE, wb, GUIDE)


GUIDE = dict(
    file=BASE + "-GUIDE.docx",
    cover=("Impact and Ease Prioritisation", "A short method and a scoring workbook for choosing which recommendations to push first", "Guide to the Excel workbook · ₹199"),
    inside=[
        "**Part A.** Why a report with twenty recommendations is a report with none.",
        "**Part B.** The method in five steps: list, score, discuss, group and agree.",
        "**Part C.** Reference assets: scoring anchors, a workshop routine and a glossary.",
        "**Part D.** A worked example: eight recommendations from a maternal and newborn health evaluation.",
    ],
    intro="This guide goes with the Excel workbook. Run the scoring with the client in the room. A ranking that the client helped to make will be used. One that it did not will be filed.",
    parts=[
        ("PART A", "Why prioritise", [
            ("body", "A report with twenty recommendations leaves the reader to choose. The reader will choose what is easiest or what the reader already wanted, and the finding that mattered most may not be on the list of things to do. A short, ranked list that the client helped to make is the nearest thing to a guarantee that something happens."),
            ("body", "The sheet scores each recommendation on three things: impact, ease and evidence. Impact and ease decide the group. Evidence breaks ties, and reminds everyone that a recommendation the findings barely support should not be first."),
            ("body", "Three rules apply:"),
            ("bullets", [
                "Score in a group of two or three, with at least one person from the client.",
                "Write down why you gave every 1 and every 5.",
                "Say in the report that the weights are a choice, and say what they are.",
            ]),
        ]),
        ("PART B", "The method", [
            ("h", "1.  List"),
            ("purpose", "Put every recommendation on one sheet."),
            ("guide", "Copy the recommendations from the draft report. Merge any that overlap, so that each is one action someone could do."),
            ("good", "Each recommendation is a verb and an object with an owner."),
            ("mistakes", "Keeping 'improve quality' as a recommendation. Nobody can do it on Monday."),
            ("h", "2.  Score"),
            ("purpose", "Make the judgement visible."),
            ("guide", "Score impact, ease and evidence from 1 to 5, using the anchors in Part C. Do it alone first and then compare."),
            ("good", "Scores differ between people, and the differences are discussed."),
            ("mistakes", "Averaging the scores before anyone has explained them."),
            ("h", "3.  Discuss"),
            ("purpose", "Find out what each person knows."),
            ("guide", "Spend the time on the scores that differ by two points or more. The person with a 5 and the person with a 2 each know something the other does not."),
            ("good", "At least one score changes after the discussion."),
            ("mistakes", "Letting the most senior person set the score."),
            ("h", "4.  Group"),
            ("purpose", "See the choice."),
            ("guide", "The sheet puts each recommendation in a group and draws the grid. Do first needs impact and ease of 4 or 5. Plan it has high impact and low ease and needs a budget or a decision. Quick win is easy and lower in impact. Park is the rest."),
            ("table", ["Group", "Impact", "Ease", "What to do"], [
                ["Do first", "4 or 5", "4 or 5", "Start now, with an owner."],
                ["Plan it", "4 or 5", "1 to 3", "Get the decision or money, and plan the steps."],
                ["Quick win", "1 to 3", "4 or 5", "Do it if it costs nothing, and do not count it as the result."],
                ["Park", "1 to 3", "1 to 3", "Leave it and say why."],
            ], [2.4, 2.2, 2.2, 9.1]),
            ("good", "No more than five recommendations are in Do first and Plan it together."),
            ("mistakes", "Calling a quick win the main result of the evaluation."),
            ("h", "5.  Agree"),
            ("purpose", "End with owners and dates."),
            ("guide", "Write an owner against every Do first and Plan it row. Put the list into the management response tracker."),
            ("good", "The client's name is against each action."),
            ("mistakes", "Finishing the meeting with a ranking and no owners."),
        ]),
        ("PART C", "Reference assets", [
            ("h", "C1.  Scoring anchors"),
            ("table", ["Score", "Impact", "Ease", "Evidence"], [
                ["5", "Fixes the main problem for most of those affected", "Within a year, with money, staff and authority in hand", "Several sources agree"],
                ["3", "Helps some of those affected", "Needs one approval or some money", "One strong source"],
                ["1", "Little effect on the problem", "Needs law, a new budget or a new institution", "Weak or single anecdotal source"],
            ], [1.4, 5.0, 5.0, 4.5]),
            ("h", "C2.  A 90-minute workshop routine"),
            ("numbered", [
                "Explain the criteria and anchors (10 minutes).",
                "Score alone (15 minutes).",
                "Compare, and discuss the scores that differ by two or more (40 minutes).",
                "Look at the groups and the grid (15 minutes).",
                "Name owners (10 minutes).",
            ]),
            ("h", "C3.  Glossary"),
            ("table", ["Term", "Meaning"], [
                ["Impact", "How much the action would improve the problem found."],
                ["Ease", "How easily the client can do it with what it has."],
                ["Evidence", "How well the evaluation's findings support the action."],
                ["Weighted score", "The sum of each score times its weight."],
            ], [3.6, 12.3]),
        ]),
        ("PART D", "Worked example: maternal and newborn health evaluation", [
            ("body", "The team scored eight recommendations with the mission's training lead and the district office. Weights were impact 40%, ease 35% and evidence 25%."),
            ("table", ["No.", "Recommendation", "Impact", "Ease", "Evidence", "Score", "Group"], [
                ["1", "New birth-asphyxia protocol for every health worker", "5", "4", "5", "4.65", "Do first"],
                ["2", "Stock oxytocin at every sub-centre", "4", "3", "4", "3.65", "Plan it"],
                ["3", "Post staff rosters", "3", "5", "3", "3.70", "Quick win"],
                ["4", "Incentive on the day of delivery", "4", "2", "3", "3.05", "Plan it"],
                ["5", "Digitise the mother-and-child records", "4", "2", "4", "3.30", "Plan it"],
                ["6", "Monthly review meetings by block", "3", "4", "3", "3.35", "Quick win"],
                ["7", "Replace delivery tables", "2", "2", "2", "2.00", "Park"],
                ["8", "Mass campaign on danger signs", "3", "3", "2", "2.75", "Park"],
            ], [0.9, 6.2, 1.5, 1.4, 1.7, 1.5, 2.7]),
            ("body", "Recommendation 1 is the only one in Do first. Recommendations 2, 4 and 5 need decisions or money, so they go to the management response tracker as Plan it. Recommendation 3 is a quick win that the block medical officers can start this month."),
            ("figures", "The scores are illustrative. A weighted score is the impact score times 0.40 plus the ease score times 0.35 plus the evidence score times 0.25."),
        ]),
    ],
)

PAGE = dict(
    slug="impact-ease-prioritisation",
    title="Impact and Ease Prioritisation",
    file=FILE,
    ext="XLSX",
    eyebrow="Template · XLSX",
    crumb="Template",
    meta="An Excel scoring workbook that ranks evaluation recommendations by impact, ease and evidence and places them on a grid, with a Word guide and a worked example. ₹199.",
    og="Rank evaluation recommendations by impact, ease and evidence, and see them on a grid. ₹199.",
    lede="A short method and a scoring workbook for choosing which recommendations to push first.",
    note="Editable XLSX workbook + a Word guide · delivered to your inbox within 24 hours of payment.",
    letter_h="Choose what to push first",
    letter=[
        "A report with twenty recommendations leaves the reader to choose, and the reader will choose what is easiest. A short, ranked list that the client helped to make is the nearest thing to a guarantee that something happens.",
        "This workbook scores each recommendation from 1 to 5 on impact, ease and evidence, applies weights you can change, ranks the list, sorts it into Do first, Plan it, Quick win or Park, and draws a grid. A Word guide gives a five-step method, scoring anchors and a 90-minute workshop routine.",
        'It uses the findings you rate with the <a href="/products/evidence-strength-rating/">Evidence Strength Rating</a> and feeds the <a href="/products/management-response-tracker/">Management Response Tracker</a>.',
    ],
    inside=[
        "A score sheet for 30 recommendations with weighted score, rank and group",
        "Weights you can change, with a check that they add to 100%",
        "A grid that places each recommendation by impact and ease",
        "A Word guide with scoring anchors and a 90-minute workshop routine",
        "A worked example: eight recommendations from a maternal and newborn health evaluation",
    ],
    sample=[
        ("Score sheet", "impact, ease and evidence, 1 to 5"),
        ("Groups", "Do first, Plan it, Quick win or Park"),
        ("Grid", "impact against ease on one page"),
        ("Scoring anchors", "what a 1, a 3 and a 5 mean"),
    ],
    sample_file=None,
    faq=[
        ("What format is it?", "An Excel workbook (.xlsx) with live formulas and dropdowns, plus a Word guide. It also opens in Google Sheets and LibreOffice."),
        ("Can I change the weights?", "Yes. Change them in the Weights sheet. Say in the report what you used."),
        ("How many recommendations can it hold?", "Thirty. Fewer than ten is better for the client."),
        ("How fast is delivery?", "Within 24 hours of confirming payment, usually much sooner."),
    ],
    search_desc="An Excel workbook that ranks evaluation recommendations by impact, ease and evidence, with a grid, a Word guide and a worked example. Buy for ₹199 via UPI.",
    search_tags=["template", "prioritisation", "recommendations", "impact ease", "consulting", "evaluation", "Excel", "premium", "paid", "₹199"],
    search_id="PROD-IMPACT-EASE",
    card="Impact and Ease Prioritisation",
    card_p="Rank recommendations and see the choice on one grid.",
    before="Stakeholder Map",
)
