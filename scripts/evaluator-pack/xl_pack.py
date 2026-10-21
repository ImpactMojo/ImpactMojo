"""Shared helpers for the Excel products: styling, and zipping the workbook with its Word guide."""
import zipfile
from pathlib import Path

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from scorecard import BORDER, HEAD_FILL, IN_FILL, NAVY, SOFT, TEAL, box, head, title  # noqa: F401
from spec_build import build_doc


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = w


def dropdown(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def note(ws, cell, text, bold=False, color="5B6776", size=10):
    ws[cell] = text
    ws[cell].font = Font(name="Calibri", size=size, bold=bold, color=color)
    ws[cell].alignment = Alignment(wrap_text=False, vertical="top")


def readme(ws, heading, sub, lines):
    title(ws, heading, sub)
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    r = 4
    for t, bold in lines:
        c = ws.cell(row=r, column=2, value=t)
        c.font = Font(name="Calibri", size=10.5, bold=bold, color=NAVY if bold else "0F172A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def package(out, base, wb, guide_spec):
    """Save the workbook, build the guide, zip both. Returns the three file names."""
    out = Path(out)
    x, g, z = f"{base}.xlsx", f"{base}-GUIDE.docx", f"{base}.zip"
    wb.save(out / x)
    build_doc(guide_spec, out / g)
    with zipfile.ZipFile(out / z, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(out / x, x)
        zf.write(out / g, g)
    return [x, g, z]


def grid_sheet(ws, heading, sub, cols, nrows, rows=None, calc=None, drops=None, first=5, fmts=None, freeze="A5"):
    """A register-style sheet. cols: [(header, width)]; calc: {col_index: lambda r -> formula}; drops: {col_index: [options]}."""
    from openpyxl.utils import get_column_letter as L
    title(ws, heading, sub)
    calc, drops, fmts = calc or {}, drops or {}, fmts or {}
    head(ws, first - 1, [c[0] for c in cols])
    ws.row_dimensions[first - 1].height = 44
    for j, (_, w) in enumerate(cols):
        ws.column_dimensions[L(j + 1)].width = w
    rows = rows or []
    for i in range(nrows):
        r = first + i
        data = rows[i] if i < len(rows) else []
        for j in range(len(cols)):
            if j in calc:
                cc = ws.cell(row=r, column=j + 1, value=calc[j](r))
                box(cc, SOFT, align="center")
            else:
                v = data[j] if j < len(data) else None
                cc = ws.cell(row=r, column=j + 1, value=v)
                box(cc, IN_FILL, align="top")
            if j in fmts:
                cc.number_format = fmts[j]
            cc.alignment = Alignment(wrap_text=True, vertical="top")
    for j, opts in drops.items():
        dropdown(ws, f"{L(j + 1)}{first}:{L(j + 1)}{first + nrows - 1}", opts)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = freeze
