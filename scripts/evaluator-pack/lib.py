"""House-format Word builder for the ImpactMojo product packs.

It reproduces the layout of the v2.0 masters in the Drive folder "ImpactMojo -
Product Masters": a navy cover block, an "inside this pack" box, Parts A to D, annotated
sections with a green "what good looks like" box and an amber "common mistakes" box,
navy-headed data tables, and a footer with page numbers. Measured from the
ImpactMojo-MEL-Plan-Pack master on 2026-10-21 (Calibri, A4, 2.54 cm margins).
"""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY, TEAL, GREY, INK = "152A49", "0C8A89", "5B6776", "0F172A"
GREEN_T, AMBER_T = "14703F", "8A6100"
FONT = "Calibri"
WIDTH_CM = 15.9


def rgb(h):
    return RGBColor.from_string(h)


def _run(p, text, size=10.5, bold=False, color=INK, italic=False, highlight=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = rgb(color)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=100, left=160, bottom=100, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _table_borders(table, color, left_sz=4, sz=4, none=False):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if none:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(left_sz if edge == "left" else sz))
            el.set(qn("w:color"), color)
        el.set(qn("w:space"), "0")
        b.append(el)
    tblPr.append(b)


def _fixed(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    # LibreOffice and Google Docs size columns from the grid, Word from the cells.
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths_cm):
            gc.set(qn("w:w"), str(int(widths_cm[i] / 2.54 * 1440)))


def _para_space(p, before=0, after=4, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _field(run, instr):
    for t, text in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if t == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        run._r.append(el)


class Pack:
    def __init__(self):
        self.d = Document()
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
        st = self.d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        fp = s.footer.paragraphs[0]
        _run(fp, "© 2026 ImpactMojo · CC BY-NC-ND 4.0 · impactmojo.in        Page ", 7.5, color=GREY)
        _field(_run(fp, "", 7.5, color=GREY), "PAGE")
        _run(fp, " of ", 7.5, color=GREY)
        _field(_run(fp, "", 7.5, color=GREY), "NUMPAGES")

    # ---------- blocks ----------
    def cover(self, title, subtitle, line3):
        t = self.d.add_table(rows=1, cols=1)
        _fixed(t, [WIDTH_CM])
        _table_borders(t, "000000", none=True)
        c = t.rows[0].cells[0]
        _shade(c, NAVY)
        _cell_margins(c, 260, 320, 260, 320)
        p = c.paragraphs[0]
        _run(p, "ImpactMojo  ·  Evaluator's Desk", 8.5, color="AEBDCC")
        _para_space(p, 0, 6)
        p = c.add_paragraph()
        _run(p, title, 20, True, "FFFFFF")
        _para_space(p, 0, 4)
        p = c.add_paragraph()
        _run(p, subtitle, 11, color="0DBEB6")
        _para_space(p, 0, 4)
        p = c.add_paragraph()
        _run(p, line3, 8.5, color="AEBDCC")
        _para_space(p, 0, 0)

    def version(self, text):
        p = self.d.add_paragraph()
        _run(p, text, 9, color=GREY)
        _para_space(p, 6, 6)

    def box(self, label, paragraphs, border=TEAL, fill="EAF6F5", label_color=TEAL):
        t = self.d.add_table(rows=1, cols=1)
        _fixed(t, [WIDTH_CM])
        _table_borders(t, border, left_sz=24, sz=4)
        c = t.rows[0].cells[0]
        _shade(c, fill)
        _cell_margins(c)
        p = c.paragraphs[0]
        _run(p, label, 8.5, True, label_color)
        _para_space(p, 0, 3)
        for text in paragraphs:
            p = c.add_paragraph()
            self._rich(p, text, 9.5)
            _para_space(p, 0, 3)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def good(self, text):
        self.box("WHAT GOOD LOOKS LIKE", [text], "1FAE63", "EAF7EF", GREEN_T)

    def mistakes(self, text):
        self.box("COMMON MISTAKES", [text], "E0A800", "FFF7E6", AMBER_T)

    def legal(self, text):
        """The law as read on the build date, stated with its source."""
        self.box("LAW AS CHECKED ON 21 OCTOBER 2026", [text], "3B6EA5", "EEF3F9", NAVY)

    def figures(self, text):
        self.box("ABOUT THE FIGURES", [text], "8A94A3", "F2F4F7", GREY)

    def part(self, label, title):
        p = self.d.add_paragraph()
        _run(p, label, 9, True, TEAL)
        _para_space(p, 14, 0)
        p.paragraph_format.keep_with_next = True
        p = self.d.add_paragraph()
        _run(p, title, 16, True, NAVY)
        _para_space(p, 0, 6)
        p.paragraph_format.keep_with_next = True

    def h(self, text):
        p = self.d.add_paragraph()
        _run(p, text, 11.5, True, NAVY)
        _para_space(p, 10, 2)
        p.paragraph_format.keep_with_next = True

    def sub(self, text):
        p = self.d.add_paragraph()
        _run(p, text, 10.5, True, NAVY)
        _para_space(p, 6, 2)
        p.paragraph_format.keep_with_next = True

    def purpose(self, text):
        p = self.d.add_paragraph()
        _run(p, "What this section is for.  ", 10, True, TEAL)
        _run(p, text, 10, color=INK)
        _para_space(p, 0, 3)

    def guide(self, text):
        p = self.d.add_paragraph()
        _run(p, text, 9.5, color=GREY)
        _para_space(p, 0, 4)

    def body(self, text):
        p = self.d.add_paragraph()
        self._rich(p, text, 10.5)
        _para_space(p, 0, 5)

    def bullets(self, items, size=10.5):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet")
            self._rich(p, it, size)
            _para_space(p, 0, 2)

    def numbered(self, items, size=10.5):
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            _run(p, f"{i}.  ", size, True, NAVY)
            self._rich(p, it, size)
            _para_space(p, 0, 2)

    def lines(self, n=2):
        for _ in range(n):
            p = self.d.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:space"), "1")
            b.set(qn("w:color"), "C9D2DD")
            bd.append(b)
            pPr.append(bd)
            _para_space(p, 0, 8)

    def table(self, header, rows, widths_cm, size=9):
        t = self.d.add_table(rows=1 + len(rows), cols=len(header))
        _fixed(t, widths_cm)
        _table_borders(t, "C9D2DD")
        for i, htxt in enumerate(header):
            c = t.rows[0].cells[i]
            _shade(c, NAVY)
            _cell_margins(c, 70, 100, 70, 100)
            _run(c.paragraphs[0], htxt, 9.5, True, "FFFFFF")
        for r, row in enumerate(rows, 1):
            for i, txt in enumerate(row):
                c = t.rows[r].cells[i]
                _shade(c, "F2F6F8" if r % 2 == 0 else "FFFFFF")
                _cell_margins(c, 60, 100, 60, 100)
                p = c.paragraphs[0]
                self._rich(p, txt, size)
        # repeat header row on page breaks
        trPr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trPr.append(th)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def page_break(self):
        self.d.add_page_break()

    def _rich(self, p, text, size):
        """**bold** inline markers only."""
        parts = text.split("**")
        for i, part in enumerate(parts):
            if part:
                _run(p, part, size, bold=(i % 2 == 1), color=INK)

    def save(self, path):
        self.d.save(path)
