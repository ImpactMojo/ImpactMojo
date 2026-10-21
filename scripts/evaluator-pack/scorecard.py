"""Bid or No-Bid Scorecard: an Excel workbook with live formulas and a Word guide.

Ships as a zip (xlsx + guide docx), the same way the other five Excel templates do.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

from lib import Pack

FILE_XLSX = "ImpactMojo-Bid-No-Bid-Scorecard.xlsx"
FILE_GUIDE = "ImpactMojo-Bid-No-Bid-Scorecard-GUIDE.docx"
FILE_ZIP = "ImpactMojo-Bid-No-Bid-Scorecard.zip"

NAVY, TEAL, INPUT, GRID = "152A49", "0C8A89", "FFF7D6", "C9D2DD"

CRITERIA = [
    # name, what to look for, weight, question to ask if the score is 1 or 2
    ("Fit with your sector and method experience",
     "The sector, population, methods and language the work needs, compared with what you and your partners have done.",
     15, "Which parts of the work sit outside our experience, and can a named partner cover them?"),
    ("The decision and the intended use are clear",
     "The ToR names the decision the evidence will inform, who makes it and when.",
     15, "Which decision will this evidence inform, who makes it, and by when?"),
    ("The questions are few and answerable",
     "A short list of questions, each specific about change, people and period, with data that could answer it.",
     10, "Which questions matter most for the decision, and which can be merged or dropped?"),
    ("The method and standard of evidence match the questions",
     "A claim of attribution comes with a comparison group and a baseline. Otherwise the ToR accepts contribution.",
     15, "Does the client expect attribution? If so, is there a comparison group and baseline data? If not, will the client accept a contribution design?"),
    ("The budget matches the work described",
     "A stated envelope that covers the sample, the team days, field costs and taxes the ToR implies.",
     15, "What is the budget envelope, and does it include field costs, expenses and taxes?"),
    ("The timeline allows a real inception and review",
     "Weeks for inception, fieldwork, analysis and client comments, with review periods written in.",
     10, "Can the inception run for two to three weeks, and can comment periods be written into the timeline?"),
    ("Payment terms and contract risk are acceptable",
     "Payments tied to submission, a fixed comment period, limits on extra tasks, findings left to the evaluator.",
     10, "Can payments be tied to submission, with a fixed comment period and a deemed-acceptance rule?"),
    ("Relationship and chance of winning",
     "Whether you know the commissioner, who sits on the panel, whether there is an incumbent, and whether the ToR was written around one provider.",
     10, "Who is on the panel, is there an incumbent, and was the ToR written around a particular provider?"),
]

KNOCKOUTS = [
    "A conflict of interest with the programme or its implementer that cannot be declared and managed.",
    "The ToR lets the client change the findings.",
    "You cannot name a qualified team leader with the days the work needs.",
    "The proposal deadline leaves you less time than you need to write a defensible bid.",
]

thin = Side(style="thin", color=GRID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
IN_FILL = PatternFill("solid", fgColor=INPUT)
SOFT = PatternFill("solid", fgColor="F2F6F8")


def head(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER


def box(c, fill=None, wrap=True, bold=False, size=10, color="0F172A", align="top"):
    c.border = BORDER
    c.alignment = Alignment(wrap_text=wrap, vertical=align)
    c.font = Font(name="Calibri", size=size, bold=bold, color=color)
    if fill:
        c.fill = fill


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", size=10, color="5B6776")
    ws.sheet_view.showGridLines = False


def scorecard_sheet(ws, example=None):
    title(ws, "Bid or no-bid scorecard", "Score the ToR before you read it twice. Yellow cells are yours to fill. Everything else calculates.")
    ex = example or {}
    ws["A4"], ws["D4"] = "Assignment", "Client"
    ws["A5"], ws["D5"] = "Date", "Scored by"
    for a in ("A4", "D4", "A5", "D5"):
        ws[a].font = Font(bold=True, name="Calibri", size=10)
    ws["B4"], ws["E4"] = ex.get("assignment", ""), ex.get("client", "")
    ws["B5"], ws["E5"] = ex.get("date", ""), ex.get("by", "")
    for a in ("B4", "E4", "B5", "E5"):
        box(ws[a], IN_FILL)
    ws.merge_cells("B4:C4"); ws.merge_cells("E4:G4"); ws.merge_cells("B5:C5"); ws.merge_cells("E5:G5")

    ws["A7"] = "Set your thresholds before you look at the ToR."
    ws["A7"].font = Font(bold=True, name="Calibri", size=10, color=TEAL)
    ws["A8"], ws["A9"] = "Go if the total is at least", "Go with conditions if the total is at least"
    ws["D8"], ws["D9"] = ex.get("go", 75), ex.get("cond", 55)
    for a in ("D8", "D9"):
        box(ws[a], IN_FILL, align="center")
    for a in ("A8", "A9"):
        ws[a].font = Font(name="Calibri", size=10)

    ws["A11"] = "Knockout checks. Any Yes means no-bid."
    ws["A11"].font = Font(bold=True, name="Calibri", size=10, color=TEAL)
    head(ws, 12, ["#", "Check", "", "", "", "", "Yes or No"])
    ws.merge_cells("B12:F12")
    for i, k in enumerate(KNOCKOUTS):
        r = 13 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value=k)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.cell(row=r, column=7, value=(ex.get("ko", ["", "", "", ""])[i]))
        box(ws.cell(row=r, column=1)); box(ws.cell(row=r, column=2)); box(ws.cell(row=r, column=7), IN_FILL, align="center")
        ws.row_dimensions[r].height = 30
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G13:G16")

    ws["A18"] = "Score each criterion from 1 (poor) to 5 (strong), and write one line of evidence from the ToR."
    ws["A18"].font = Font(bold=True, name="Calibri", size=10, color=TEAL)
    head(ws, 19, ["#", "Criterion", "What to look for", "Weight", "Score 1 to 5", "Why this score", "Points", "Question to ask the client if the score is 1 or 2"])
    scores = ex.get("scores", [None] * 8)
    whys = ex.get("why", [""] * 8)
    for i, (name, look, w, q) in enumerate(CRITERIA):
        r = 20 + i
        vals = [i + 1, name, look, w, scores[i], whys[i], f'=IF(E{r}="","",D{r}*E{r}/5)', f'=IF(E{r}="","",IF(E{r}<=2,"ASK: {q}",""))']
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v)
            box(c, IN_FILL if j in (5, 6) else None, align="center" if j in (1, 4, 5, 7) else "top")
        ws.row_dimensions[r].height = 62
    dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                         errorTitle="Score", error="Enter a whole number from 1 to 5.")
    ws.add_data_validation(dv2)
    dv2.add("E20:E27")

    ws["A29"], ws["D29"] = "Weights add to", "=SUM(D20:D27)"
    ws["A30"], ws["D30"] = "Total score out of 100", '=IF(COUNT(E20:E27)<8,"",SUM(G20:G27))'
    ws["A31"], ws["D31"] = "Lowest single score", '=IF(COUNT(E20:E27)<8,"",MIN(E20:E27))'
    ws["A32"], ws["D32"] = "Criteria scored 1 (each is a condition to resolve)", '=COUNTIF(E20:E27,1)'
    ws["A33"], ws["D33"] = "Knockouts answered Yes", '=COUNTIF(G13:G16,"Yes")'
    for r in range(29, 34):
        ws[f"A{r}"].font = Font(name="Calibri", size=10, bold=True)
        box(ws[f"D{r}"], SOFT, align="center", bold=True)
    ws["A35"] = "DECISION"
    ws["A35"].font = Font(name="Calibri", size=12, bold=True, color=NAVY)
    ws["B35"] = ('=IF(COUNT(E20:E27)<8,"Score all eight criteria first",'
                 'IF(D33>0,"NO-BID: a knockout applies",'
                 'IF(D30>=D8,IF(D31=1,"GO WITH CONDITIONS","GO"),'
                 'IF(D30>=D9,"GO WITH CONDITIONS","NO-BID"))))')
    ws.merge_cells("B35:E35")
    ws["B35"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws["B35"].fill = PatternFill("solid", fgColor=NAVY)
    ws["B35"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[35].height = 28
    ws["A37"] = "A bid that scores 1 on any criterion becomes a condition: raise it as a clarification question before you commit days to the proposal. Write down your reason for the decision on the Bid log sheet."
    ws["A37"].font = Font(name="Calibri", size=9, italic=True, color="5B6776")
    ws.merge_cells("A37:H37")
    ws.row_dimensions[37].height = 28
    ws["A37"].alignment = Alignment(wrap_text=True, vertical="top")

    red = PatternFill("solid", bgColor="F8D7DA")
    amber = PatternFill("solid", bgColor="FFF1C2")
    green = PatternFill("solid", bgColor="D5F0DF")
    ws.conditional_formatting.add("B35:E35", FormulaRule(formula=['LEFT($B$35,2)="NO"'], fill=PatternFill("solid", bgColor="B42318"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add("B35:E35", FormulaRule(formula=['LEFT($B$35,7)="GO WITH"'], fill=PatternFill("solid", bgColor="8A6100"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add("B35:E35", FormulaRule(formula=['$B$35="GO"'], fill=PatternFill("solid", bgColor="14703F"), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add("E20:E27", FormulaRule(formula=['E20=1'], fill=red))
    ws.conditional_formatting.add("E20:E27", FormulaRule(formula=['E20=2'], fill=amber))
    ws.conditional_formatting.add("E20:E27", FormulaRule(formula=['E20>=4'], fill=green))

    for col, w in zip("ABCDEFGH", (11, 30, 40, 9, 10, 44, 9, 52)):
        ws.column_dimensions[col].width = w
    ws["A4"].alignment = ws["A5"].alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "A4"


def cost_sheet(ws, ex=None):
    ex = ex or {}
    title(ws, "Cost of bidding", "A bid costs days you could bill elsewhere. Test whether the expected return covers them.")
    rows = [
        ("Days you and your team will spend writing the proposal", ex.get("days", None), "0.0"),
        ("Your day rate (₹)", ex.get("rate", None), "#,##0"),
        ("Fee if you win (₹)", ex.get("fee", None), "#,##0"),
        ("Margin on the fee if you win (%)", ex.get("margin", None), "0%"),
        ("Your honest chance of winning (%)", ex.get("p", None), "0%"),
    ]
    for i, (lab, val, fmt) in enumerate(rows):
        r = 4 + i
        ws.cell(row=r, column=1, value=lab).font = Font(name="Calibri", size=10)
        c = ws.cell(row=r, column=3, value=val)
        box(c, IN_FILL, align="center")
        c.number_format = fmt
    calc = ["Cost of the bid (₹)", "Margin if you win (₹)", "Expected value of bidding (₹)", "Win chance needed to break even"]
    # layout: rows 10-13 for results
    ws["A10"], ws["A11"], ws["A12"], ws["A13"] = calc
    ws["C10"] = '=IF(OR(C4="",C5=""),"",C4*C5)'
    ws["C11"] = '=IF(OR(C6="",C7=""),"",C6*C7)'
    ws["C12"] = '=IF(OR(C10="",C11="",C8=""),"",C8*C11-C10)'
    ws["C13"] = '=IF(OR(C10="",C11="",C11=0),"",C10/C11)'
    for r, fmt in ((10, "#,##0"), (11, "#,##0"), (12, "#,##0"), (13, "0.0%")):
        ws[f"A{r}"].font = Font(name="Calibri", size=10, bold=True)
        box(ws[f"C{r}"], SOFT, align="center", bold=True)
        ws[f"C{r}"].number_format = fmt
    ws["A15"] = '=IF(C12="","",IF(C12>0,"On money alone, this bid is worth writing.","On money alone, this bid is not worth writing. Bid only for a reason you can name: a relationship, a reference, a method you want to build."))'
    ws["A15"].font = Font(name="Calibri", size=10, bold=True, color=NAVY)
    ws.merge_cells("A15:G15")
    ws["A15"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[15].height = 34
    ws["A17"] = "Be honest about the chance of winning, and check your own record on the Bid log sheet."
    ws["A17"].font = Font(name="Calibri", size=9, italic=True, color="5B6776")
    ws.merge_cells("A17:G17")
    ws["A17"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[17].height = 28
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 4
    ws.column_dimensions["C"].width = 18


def log_sheet(ws):
    title(ws, "Bid log", "Record every decision, including the no-bids. After ten entries you will see your own pattern.")
    cols = ["Date", "Client", "Assignment", "Fee if won (₹)", "Score", "Decision", "Outcome", "Reason for the decision", "What you would do differently"]
    head(ws, 4, cols)
    for r in range(5, 35):
        for c in range(1, 10):
            box(ws.cell(row=r, column=c), IN_FILL, align="top")
        ws.cell(row=r, column=4).number_format = "#,##0"
    dv = DataValidation(type="list", formula1='"GO,GO WITH CONDITIONS,NO-BID"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Won,Lost,Withdrawn,Not bid,Waiting"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dv2)
    dv.add("F5:F34"); dv2.add("G5:G34")
    ws["K4"], ws["K5"], ws["K6"], ws["K7"], ws["K8"] = "Your record", "Bids submitted", "Bids won", "Win rate", "No-bids recorded"
    ws["L5"] = '=COUNTIF(G5:G34,"Won")+COUNTIF(G5:G34,"Lost")'
    ws["L6"] = '=COUNTIF(G5:G34,"Won")'
    ws["L7"] = '=IF(L5=0,"",L6/L5)'
    ws["L8"] = '=COUNTIF(F5:F34,"NO-BID")'
    ws["K4"].font = Font(bold=True, color=NAVY, name="Calibri", size=11)
    for r in range(5, 9):
        ws[f"K{r}"].font = Font(name="Calibri", size=10)
        box(ws[f"L{r}"], SOFT, align="center", bold=True)
    ws["L7"].number_format = "0%"
    for col, w in zip("ABCDEFGHI", (12, 22, 30, 16, 8, 22, 12, 40, 40)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["K"].width = 20
    ws.column_dimensions["L"].width = 12
    ws.freeze_panes = "A5"


def readme_sheet(ws):
    title(ws, "ImpactMojo · Bid or no-bid scorecard", "Version 1.0 · October 2026 · Part of the Evaluation Essentials Kit")
    lines = [
        ("How to use it", ""),
        ("1", "Open the Scorecard sheet. Set your go and conditions thresholds before you read the ToR twice."),
        ("2", "Answer the four knockout checks. Any Yes means no-bid."),
        ("3", "Score the eight criteria from 1 to 5 and write one line of evidence from the ToR for each."),
        ("4", "Read the decision. A score of 1 on any criterion becomes a condition: the sheet writes the clarification question for you."),
        ("5", "Open Cost of bidding and test whether the expected return covers the days the proposal will take."),
        ("6", "Record the decision, and your reason, on the Bid log sheet. Review the log every ten bids."),
        ("", ""),
        ("Colours", "Yellow cells are inputs. Grey cells are calculated. Do not type over a grey cell."),
        ("The Word guide", "Explains how to score each criterion, lists the red-flag clauses to look for in a ToR, and walks through a worked example."),
        ("Worked example sheet", "A completed scorecard for the maternal and newborn health case used across the ImpactMojo kit. Compare it with your own."),
        ("", ""),
        ("Licence", "© 2026 ImpactMojo · CC BY-NC-ND 4.0 · impactmojo.in. Use it in your own practice. Do not resell or redistribute it."),
    ]
    for i, (a, b) in enumerate(lines):
        r = 4 + i
        ws.cell(row=r, column=1, value=a).font = Font(bold=True, name="Calibri", size=10, color=TEAL if i == 0 else NAVY)
        c = ws.cell(row=r, column=2, value=b)
        c.font = Font(name="Calibri", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30 if len(b) > 90 else 18
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100


EXAMPLE = dict(
    assignment="Outcome evaluation, Surakshit Maa programme (first draft ToR)",
    client="Funding foundation",
    date="Week 0",
    by="Team leader",
    go=75,
    cond=55,
    ko=["No", "No", "No", "No"],
    scores=[5, 3, 3, 1, 3, 2, 2, 3],
    why=[
        "Maternal and newborn health, mixed methods, field language covered by the team.",
        "ToR says 'to inform the funder's decision on phase 2' with no date and no named decision-maker.",
        "Six questions. One asks about maternal mortality, which no feasible sample can detect.",
        "ToR asks for 'impact and attribution' with no comparison group and no baseline survey.",
        "₹30 lakh stated. The work described is closer to ₹38 to 42 lakh, so scope must be trimmed.",
        "10 weeks including fieldwork across three blocks, with no room for client comments.",
        "50% of the fee on approval of the final report. 'Any other tasks' clause present.",
        "New funder. No known incumbent. Panel not stated.",
    ],
)
EXAMPLE_COST = dict(days=5, rate=25000, fee=3000000, margin=0.25, p=0.30)


def build_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    readme_sheet(ws)
    scorecard_sheet(wb.create_sheet("Scorecard"))
    cost_sheet(wb.create_sheet("Cost of bidding"))
    log_sheet(wb.create_sheet("Bid log"))
    scorecard_sheet(wb.create_sheet("Worked example"), EXAMPLE)
    cost_sheet(wb.create_sheet("Worked cost"), EXAMPLE_COST)
    wb["Worked example"].title = "Worked example"
    wb.save(path)


# ------------------------------------------------------------------ guide ----

def build_guide(path):
    k = Pack()
    k.cover(
        "Bid or No-Bid Scorecard",
        "The complete pack: scoring workbook, anchored criteria, red-flag clause reader and a worked example",
        "Template · XLSX + Word guide · ₹199",
    )
    k.version("Version 1.0   ·   October 2026   ·   Part of the Evaluation Essentials Kit")
    k.box("WHAT IS INSIDE THIS PACK", [
        "**The workbook.** A scorecard with live formulas, a cost-of-bidding test, a bid log and a completed worked example.",
        "**Part A.** Why a gate, and what it protects.",
        "**Part B.** How to use the workbook, and what a 1, a 3 and a 5 look like for each criterion.",
        "**Part C.** Reference assets: red-flag clauses to look for in a ToR, with the question to ask for each, and rules for when to bid anyway.",
        "**Part D.** A worked example: scoring the first draft ToR for the maternal and newborn health programme.",
    ])
    k.guide("Write down your thresholds before you open the ToR. A threshold set after you have seen the score is a rationalisation. Yellow [LEGAL CHECK] notes mark points where a statute, tender rule or contract changes the answer.")

    k.part("PART A", "Why a gate")
    k.body("A serious proposal takes days of unpaid work. Most of the assignments that cost an evaluator money were visible in the ToR: a question set nobody could answer, a budget half the size of the work, a payment clause that left the fieldwork unfunded. A short, repeatable gate moves those signals from a feeling to a number and turns each one into a question you can put to the client before you commit days.")
    k.body("The gate does three jobs. It saves the days you would have spent on a bid you could not win or could not afford to deliver. It improves the bids you do write, because the low scores become clarification questions. And it builds a record, so that after ten bids you can see which kinds of work you win and which you only think you will.")

    k.part("PART B", "Using the workbook")
    k.h("1.  Set the thresholds, then run the knockouts")
    k.body("The default thresholds are 75 for a go and 55 for a go with conditions. Change them to suit your practice, but change them before you score. Then answer the four knockout checks. A conflict of interest you cannot declare, a clause that lets the client change findings, no qualified team leader for the days needed, or a deadline that leaves too little time each end the discussion.")
    k.h("2.  Score the criteria")
    k.guide("Weights add to 100. A score of 5 earns the full weight, a 3 earns 60% of it, a 1 earns 20%. Write one line of evidence from the ToR beside each score.")
    k.table(["Criterion (weight)", "A 1 looks like", "A 3 looks like", "A 5 looks like"], [
        ["Fit with your sector and method experience (15)", "Different sector, method and language, with nobody on your team who has done it.", "Adjacent sector or method. You would need a partner for one piece.", "Work you have done, in the same sector, method and language, with a reference."],
        ["The decision and intended use are clear (15)", "'To assess the programme.' No decision, no date.", "A decision is named but not its owner or date.", "Names the decision, the decision-maker and the date."],
        ["The questions are few and answerable (10)", "Over fifteen questions, or questions no sample could answer.", "Six to ten questions, one or two unanswerable.", "Up to six specific questions, each with data that could answer it."],
        ["The method and standard of evidence match the questions (15)", "Asks for attribution with no comparison group or baseline.", "Method is open, but the ToR's claims run ahead of any design.", "The method is left to bidders within constraints, and claims match the design."],
        ["The budget matches the work described (15)", "Envelope under half of what the work needs, or not stated at all.", "Envelope 20 to 30% short, with scope you could trim.", "Envelope covers the work, with field costs and taxes stated."],
        ["The timeline allows a real inception and review (10)", "Fewer weeks than the fieldwork needs, with no review time.", "Tight, but an inception and one review round fit.", "Weeks for inception, fieldwork, analysis and client comment periods."],
        ["Payment terms and contract risk are acceptable (10)", "Most of the fee on final approval, open-ended extra tasks, findings approved by the client.", "Some payment on signing, with one clause to renegotiate.", "Payments on submission, fixed comment period, findings left to the evaluator."],
        ["Relationship and chance of winning (10)", "ToR appears written around another provider.", "Unknown panel, no incumbent evident.", "You know the commissioner, and the panel values your approach."],
    ], [4.4, 4.0, 3.8, 3.7], size=8.5)
    k.h("3.  Read the decision, and the conditions")
    k.body("The sheet reports GO, GO WITH CONDITIONS or NO-BID. A score of 1 on any criterion downgrades a go to a go with conditions, because a single fatal weakness can undo a high total. The sheet also writes the clarification question for every criterion scored 1 or 2. Send those questions to the client before the deadline for queries. How the client answers is itself information: a client who will not clarify the design is telling you something.")
    k.h("4.  Test the economics")
    k.body("Open the Cost of bidding sheet. Enter the days the proposal will take, your day rate, the fee, your margin and your honest chance of winning. The sheet reports the cost of the bid, the margin if you win, the expected value, and the win chance you need to break even. A bid with negative expected value can still be worth writing. The sheet asks you to name the reason.")
    k.h("5.  Keep the log")
    k.body("Record every decision on the Bid log sheet, including the no-bids and the reason. Review it every ten entries: which scores predicted a win, which conditions the client accepted, and which kinds of work you keep losing.")

    k.part("PART C", "Reference assets")
    k.h("C1.  Red-flag clauses in a ToR, and the question to ask")
    k.table(["Clause or signal", "Why it matters", "Question to ask"], [
        ["'The consultant will also undertake any other tasks assigned by the programme team.'", "An open-ended scope on a fixed fee.", "Can extra tasks be agreed through a written change request with days and fee?"],
        ["'100% of the fee on approval of the final report.'", "You carry every fieldwork cost until a person you do not control signs off.", "Can payments be tied to submission of each deliverable?"],
        ["'Findings will be finalised after approval by the programme team.'", "The programme becomes the judge of its own evaluation.", "Can the client correct facts while the evaluator decides the findings?"],
        ["Impact or attribution requested with no comparison group or baseline.", "The design cannot support the claim.", "Will the client accept a contribution design, or is there baseline and comparison data?"],
        ["More than ten evaluation questions, or a 'full list in Annex 2'.", "A long list yields a thin answer to each.", "Which questions matter most for the decision?"],
        ["A timeline too short for an inception and fieldwork.", "No time to pilot tools or train enumerators.", "Can the inception phase run for two to three weeks?"],
        ["'A high-quality evaluation report' as the only deliverable.", "Nothing defines done.", "What length, format and review rounds does the client expect?"],
        ["No budget stated.", "You cannot tell whether the work is affordable.", "What is the budget envelope, and does it include taxes and field costs?"],
        ["The client owns all tools and data, including your pre-existing methods.", "You lose your own instruments.", "Can pre-existing tools stay with the evaluator under a licence to the client?"],
        ["Eligibility rules, such as a minimum number of similar assignments or turnover.", "You may be ineligible before you start.", "Check the eligibility rules before you spend any days."],
        ["Earnest money, performance security or a penalty clause in a tender.", "A cost and a risk that sits outside the fee.", "What are the amounts, and when are they returned?"],
    ], [5.2, 5.0, 5.7], size=8.5)
    k.legal("For a public tender, read the tender document and the procuring body's rules on eligibility, earnest money and performance security before you score. This pack does not state their amounts or conditions.")
    k.h("C2.  When to bid anyway")
    k.body("A bid that scores no-bid or a weak go can still be worth writing, if you can name the reason and cap the cost. Three reasons hold up: the client is a relationship you want for the next ten years, the assignment gives you a reference in a sector you are entering, or you will learn a method you intend to offer. Write the reason in the Bid log, set a ceiling on the days, and decide in advance that you will walk away if the clarification answers are poor.")
    k.h("C3.  Glossary")
    k.table(["Term", "Meaning"], [
        ["ToR", "Terms of reference: the client's description of the assignment, its purpose, questions, deliverables, timeline and budget."],
        ["Knockout", "A fact that ends the discussion whatever the score."],
        ["Condition", "A weakness you will raise with the client before you commit to the bid."],
        ["Expected value", "The chance of winning multiplied by the margin if you win, minus the cost of writing the bid."],
        ["Incumbent", "The provider already doing the work or the previous round of it."],
        ["Eligibility rules", "Requirements a bidder must meet to be considered, such as years of experience or number of similar assignments."],
    ], [3.8, 12.1])

    k.part("PART D", "Worked example: scoring a first draft ToR")
    k.box("ABOUT THIS EXAMPLE", [
        "The foundation behind the Surakshit Maa programme sends a first draft ToR for an outcome evaluation. It is the draft the programme team wrote before using a commissioning pack. The evaluator scores it. The completed workbook is on the Worked example sheet. Names and figures are illustrative.",
    ])
    k.table(["Criterion", "Score", "Evidence from the ToR"], [
        ["Fit with your sector and method experience", "5", "Maternal and newborn health, mixed methods, field languages covered."],
        ["The decision and intended use are clear", "3", "Decision named for phase 2, no date, no decision-maker."],
        ["The questions are few and answerable", "3", "Six questions, one on maternal mortality that no feasible sample can detect."],
        ["The method and standard of evidence match the questions", "1", "'Impact and attribution' with no comparison group and no baseline survey."],
        ["The budget matches the work described", "3", "₹30 lakh stated. The work is nearer ₹38 to 42 lakh, so scope must be trimmed."],
        ["The timeline allows a real inception and review", "2", "Ten weeks including fieldwork in three blocks, with no comment periods."],
        ["Payment terms and contract risk are acceptable", "2", "50% on approval of the final report, and an 'any other tasks' clause."],
        ["Relationship and chance of winning", "3", "New funder, no incumbent evident, panel not stated."],
    ], [6.0, 1.4, 8.5], size=9)
    k.body("**Result.** The weighted total is 56 out of 100: 15 + 9 + 6 + 3 + 9 + 4 + 4 + 6. That sits between the conditions threshold of 55 and the go threshold of 75, so the decision is GO WITH CONDITIONS. The score of 1 on method is the condition that matters most.")
    k.body("**The clarification questions the sheet writes.** Does the client expect attribution, and is there a comparison group and baseline data? Can the inception run for two to three weeks, and can comment periods be written into the timeline? Can payments be tied to submission, with a fixed comment period?")
    k.body("**The cost test.** Five days of the team leader at ₹25,000 is ₹1,25,000. A ₹30 lakh fee at a 25% margin is ₹7,50,000. At a 30% chance of winning the expected value is ₹1,00,000, and the bid breaks even at a win chance of about 17%. On money alone the bid is worth writing, as long as the clarification answers are good.")
    k.body("**What happened next.** The evaluator sent the three questions. The foundation revised the ToR: a contribution design with a comparison, a 13-week timeline, a budget of up to ₹38 to 42 lakh, and payments tied to deliverables. That revised ToR is the filled example in the ImpactMojo Terms of Reference pack.")
    k.save(path)


if __name__ == "__main__":
    import sys
    build_xlsx(sys.argv[1])
    build_guide(sys.argv[2])
